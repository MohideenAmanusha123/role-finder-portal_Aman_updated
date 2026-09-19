"""
resume_editor.py
----------------
Conservatively edits a resume's Skills and Summary sections using only
skills/applications the user confirms. Supports a structured preview/diff.
"""

import io
import json
import os
import re
from html import escape
from urllib import error as urllib_error
from urllib import request as urllib_request

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, ListFlowable, ListItem
from docx import Document
from docx.shared import Pt, RGBColor

from roles_data import SKILL_APPLICATIONS

INK = colors.HexColor("#20201C")
INK_SOFT = colors.HexColor("#65605A")
GOLD = colors.HexColor("#8A6D1F")
HAIRLINE = colors.HexColor("#D8D2C2")

BULLET_RE = re.compile(r"^\s*[-•*\u2022\u25CF\u2013]\s+")
SECTION_HEADER_PATTERN = re.compile(
    r"^\s*(summary|professional summary|objective|profile|experience|work experience|professional experience|"
    r"education|skills|technical skills|projects|certifications|achievements)\s*:?\s*$",
    re.IGNORECASE,
)
HEADER_MAP = {
    "summary": "Summary", "professional summary": "Summary", "objective": "Summary", "profile": "Summary",
    "experience": "Experience", "work experience": "Experience",
    "professional experience": "Experience", "education": "Education",
    "skills": "Skills", "technical skills": "Skills", "projects": "Projects",
    "certifications": "Certifications", "achievements": "Achievements",
}
ACRONYM_MAP = {
    "sql": "SQL", "aws": "AWS", "gcp": "GCP", "api": "API", "rest api": "REST API",
    "qa": "QA", "ui": "UI", "ux": "UX", "ci/cd": "CI/CD", "html": "HTML", "css": "CSS",
    "sql server": "SQL Server", "node.js": "Node.js", "power bi": "Power BI",
}


def format_skill_label(skill: str) -> str:
    return ACRONYM_MAP.get(skill.lower(), skill.title())


def _split_sections(text: str):
    lines, preamble, sections = text.splitlines(), [], []
    current_header, current_lines, started = None, [], False
    for raw_line in lines:
        stripped = raw_line.strip()
        m = SECTION_HEADER_PATTERN.match(stripped) if stripped else None
        if m:
            if not started:
                preamble = current_lines
            else:
                sections.append((current_header, current_lines))
            current_header = HEADER_MAP.get(m.group(1).lower(), m.group(1).title())
            current_lines, started = [], True
        else:
            current_lines.append(raw_line)
    if started:
        sections.append((current_header, current_lines))
    else:
        preamble = current_lines
    return preamble, sections


def _section_text(sections, header):
    for h, lines in sections:
        if h == header:
            return " ".join(l.strip() for l in lines if l.strip())
    return ""


def _skill_tokens(lines):
    tokens = set()
    for line in lines:
        cleaned = BULLET_RE.sub("", line).strip()
        for token in re.split(r"[,;|]", cleaned):
            normalized = re.sub(r"\s+", " ", token.strip()).lower()
            if normalized:
                tokens.add(normalized)
    return tokens


def rewrite_summary(original_summary: str, skills, applications, role_name=None, experience_signal=None) -> str:
    """Create polished, factual summary text without inventing experience."""
    labels = [format_skill_label(s) for s in skills]
    apps = [a for vals in (applications or {}).values() for a in vals if a]
    apps = list(dict.fromkeys(apps))
    if original_summary.strip():
        base = re.sub(r"\s+", " ", original_summary).strip()
        if base[-1:] not in ".!?":
            base += "."
        addition = f" Core strengths include {', '.join(labels)}." if labels else ""
        if role_name:
            addition = f" Aligned with {role_name} opportunities." + addition
        if apps:
            addition += f" Proficient with confirmed tools including {', '.join(apps)}."
        if experience_signal:
            addition += f" Experience profile: {experience_signal}-level."
        return (base + addition).strip()
    if not labels:
        return ""
    role_text = f" pursuing {role_name} opportunities" if role_name else ""
    text = f"Professional{role_text} with core strengths in {', '.join(labels)}."
    if apps:
        text += f" Proficient with confirmed tools including {', '.join(apps)}."
    if experience_signal:
        text += f" Brings a {experience_signal}-level experience profile."
    return text


def generate_ai_summary(original_summary, skills, applications, role_name=None, experience_signal=None):
    """Generate a summary with OpenAI by default, or local Ollama when selected."""
    provider = os.environ.get("AI_PROVIDER", "openai").strip().lower()
    if provider not in {"ollama", "openai"}:
        raise ValueError("AI_PROVIDER must be 'ollama' or 'openai'.")
    model = os.environ.get(
        "OLLAMA_MODEL" if provider == "ollama" else "OPENAI_MODEL",
        "llama3.2" if provider == "ollama" else "gpt-4o-mini",
    )
    labels = ", ".join(format_skill_label(s) for s in skills)
    apps = ", ".join(dict.fromkeys(a for vals in (applications or {}).values() for a in vals if a))
    prompt = (
        "Write a polished, natural-sounding professional resume summary in 2-4 sentences. "
        "Use clear business English, strong but accurate wording, and no first-person pronouns. "
        "Preserve the factual substance of the existing summary when one is provided. "
        "Do not invent employers, achievements, metrics, certifications, or experience. "
        "Do not mention these instructions, confirmed tools, or an AI system. Return only the summary. "
        f"Target role: {role_name or 'the target role'}. "
        f"Existing summary: {original_summary or 'None'}. "
        f"Confirmed skills: {labels or 'None'}. Confirmed tools: {apps or 'None'}. "
        f"Experience signal: {experience_signal or 'Not provided'}."
    )
    if provider == "ollama":
        endpoint = os.environ.get("OLLAMA_URL", "http://localhost:11434/api/generate")
        payload = json.dumps({
            "model": model,
            "prompt": (
                "You edit resumes conservatively. Return only the summary text, "
                "without a heading or explanation.\n\n" + prompt
            ),
            "stream": False,
            "options": {"temperature": 0.3},
        }).encode("utf-8")
        headers = {"Content-Type": "application/json"}
    else:
        api_key = os.environ.get("OPENAI_API_KEY", "").strip()
        if not api_key:
            raise ValueError("OPENAI_API_KEY is required when AI_PROVIDER=openai.")
        endpoint = os.environ.get("OPENAI_API_URL", "https://api.openai.com/v1/chat/completions")
        payload = json.dumps({
            "model": model,
            "temperature": 0.3,
            "messages": [
                {"role": "system", "content": "You edit resumes conservatively and return only the summary text."},
                {"role": "user", "content": prompt},
            ],
        }).encode("utf-8")
        headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}
    req = urllib_request.Request(endpoint, data=payload, headers=headers, method="POST")
    try:
        # Kept deliberately short: this call runs synchronously inside a
        # request handler. Under gunicorn's default sync worker model, a
        # single worker process serves one request at a time, so a slow AI
        # call here blocks that worker from serving anyone else until it
        # returns or times out. See the Procfile for the threaded-worker
        # setting that keeps this from stalling unrelated requests, and
        # app.py's tighter rate limit on the routes that reach this function.
        with urllib_request.urlopen(req, timeout=12) as response:
            result = json.loads(response.read().decode("utf-8"))
    except urllib_error.HTTPError as exc:
        if exc.code == 429:
            retry_after = exc.headers.get("Retry-After", "") if exc.headers else ""
            wait_hint = f" Retry after {retry_after} seconds." if retry_after.isdigit() else ""
            if provider == "ollama":
                message = (
                    "The local AI service is rate-limited or busy."
                    f"{wait_hint} Try again shortly or choose the standard summary option."
                )
            else:
                message = (
                    "The AI summary service rate limit or usage quota was reached."
                    f"{wait_hint} Check your provider usage/billing limits, or choose the standard summary option."
                )
            raise ValueError(message) from exc
        raise ValueError(
            f"Ollama is not running at {endpoint}. Install Ollama, run "
            f"'ollama run {model}', and try again." if provider == "ollama"
            else f"AI summary service could not be reached: {exc}"
        ) from exc
    except (urllib_error.URLError, TimeoutError) as exc:
        raise ValueError(
            f"Ollama is not running at {endpoint}. Install Ollama, run "
            f"'ollama run {model}', and try again." if provider == "ollama"
            else f"AI summary service could not be reached: {exc}"
        ) from exc
    try:
        summary = (
            result["response"].strip()
            if provider == "ollama"
            else result["choices"][0]["message"]["content"].strip()
        )
    except (KeyError, IndexError, TypeError, AttributeError) as exc:
        raise ValueError("AI summary service returned an invalid response.") from exc
    if not summary:
        raise ValueError("AI summary service returned an empty summary.")
    return summary


def build_edited_resume(
    original_text: str,
    matched_skills,
    confirmed_skills,
    applications=None,
    summary_mode="update",
    role_name=None,
    experience_signal=None,
) -> dict:
    preamble, sections = _split_sections(original_text)
    changes = []
    applications = applications or {}
    all_skills = sorted(set(matched_skills or []) | set(confirmed_skills or []))
    skills_line = ", ".join(format_skill_label(s) for s in all_skills) if all_skills else ""

    skills_idx = next((i for i, (h, _) in enumerate(sections) if h == "Skills"), None)
    if skills_idx is not None:
        existing_lines = sections[skills_idx][1]
        existing_tokens = _skill_tokens(existing_lines)
        new_labels = [
            format_skill_label(skill)
            for skill in all_skills
            if format_skill_label(skill).lower() not in existing_tokens
        ]
        if new_labels:
            sections[skills_idx] = ("Skills", existing_lines + [", ".join(new_labels)])
            changes.append("Added matched and user-confirmed skills while preserving existing skills.")
    elif skills_line:
        insert_at = next((i for i, (h, _) in enumerate(sections) if h == "Experience"), None)
        if insert_at is None:
            insert_at = next((i for i, (h, _) in enumerate(sections) if h == "Education"), None)
        sections.insert(insert_at if insert_at is not None else len(sections), ("Skills", [skills_line]))
        changes.append("Added a new Skills section.")

    if confirmed_skills:
        labels = ", ".join(format_skill_label(s) for s in sorted(confirmed_skills))
        changes.append(f"Included skills you confirmed you know: {labels}.")

    original_summary = _section_text(_split_sections(original_text)[1], "Summary")
    if summary_mode == "keep":
        new_summary = original_summary
    elif summary_mode == "ai":
        new_summary = generate_ai_summary(original_summary, all_skills, applications, role_name, experience_signal)
    elif summary_mode == "update":
        new_summary = rewrite_summary(original_summary, all_skills, applications, role_name, experience_signal)
    else:
        raise ValueError("Summary mode must be keep, update, or ai.")
    summary_idx = next((i for i, (h, _) in enumerate(sections) if h == "Summary"), None)
    if new_summary:
        if summary_idx is not None:
            sections[summary_idx] = ("Summary", [new_summary])
            changes.append(
                "Updated the Summary for the target role and confirmed skills."
                if summary_mode != "keep" else "Kept the existing Summary unchanged."
            )
        else:
            insert_at = 0
            sections.insert(insert_at, ("Summary", [new_summary]))
            changes.append("Added a Summary based only on confirmed skills and applications.")

    diff = {
        "skills_added": [format_skill_label(s) for s in sorted(set(confirmed_skills or []))],
        "applications": applications,
        "summary_before": original_summary,
        "summary_after": new_summary,
        "changes": changes,
    }
    return {"preamble": preamble, "sections": sections, "changes": changes, "diff": diff}


def render_edited_resume_pdf(edited: dict) -> io.BytesIO:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=.75*inch, rightMargin=.75*inch, topMargin=.7*inch, bottomMargin=.7*inch)
    styles = {
        "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=18, textColor=INK, leading=22, spaceAfter=6),
        "contact": ParagraphStyle("contact", fontName="Helvetica", fontSize=9.5, textColor=INK_SOFT, spaceAfter=1),
        "header": ParagraphStyle("header", fontName="Helvetica-Bold", fontSize=11.5, textColor=GOLD, spaceBefore=14, spaceAfter=3),
        "body": ParagraphStyle("body", fontName="Helvetica", fontSize=10, textColor=INK, leading=14, spaceAfter=3),
    }
    story = []
    preamble_lines = [l.strip() for l in edited["preamble"] if l.strip()]
    if preamble_lines:
        story.append(Paragraph(escape(preamble_lines[0]), styles["name"]))
        for l in preamble_lines[1:]:
            story.append(Paragraph(escape(l), styles["contact"]))
        story.append(Spacer(1, 10))
    for header, lines in edited["sections"]:
        content = [l.strip() for l in lines if l.strip()]
        if not content:
            continue
        story += [Paragraph(escape(header.upper()), styles["header"]), HRFlowable(width="100%", thickness=.75, color=HAIRLINE, spaceAfter=6)]
        bullets = []
        for l in content:
            if BULLET_RE.match(l):
                bullets.append(ListItem(Paragraph(escape(BULLET_RE.sub("", l)), styles["body"]), spaceAfter=3))
            else:
                if bullets:
                    story.append(ListFlowable(bullets, bulletType="bullet", start="-")); bullets = []
                story.append(Paragraph(escape(l), styles["body"]))
        if bullets:
            story.append(ListFlowable(bullets, bulletType="bullet", start="-"))
        story.append(Spacer(1, 6))
    doc.build(story); buf.seek(0); return buf


def render_edited_resume_docx(edited: dict) -> io.BytesIO:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    preamble_lines = [l.strip() for l in edited["preamble"] if l.strip()]
    if preamble_lines:
        p = doc.add_paragraph(); run = p.add_run(preamble_lines[0]); run.bold = True; run.font.size = Pt(18)
        for l in preamble_lines[1:]:
            doc.add_paragraph(l)
        doc.add_paragraph()
    gold = RGBColor(0x8A, 0x6D, 0x1F)
    for header, lines in edited["sections"]:
        content = [l.strip() for l in lines if l.strip()]
        if not content: continue
        h = doc.add_heading(level=2); r = h.add_run(header.upper()); r.font.color.rgb = gold; r.font.size = Pt(12)
        for l in content:
            doc.add_paragraph(BULLET_RE.sub("", l), style="List Bullet" if BULLET_RE.match(l) else None)
    buf = io.BytesIO(); doc.save(buf); buf.seek(0); return buf
